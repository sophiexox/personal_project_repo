from pathlib import Path
import re

import pandas as pd
from openpyxl import load_workbook

# file paths
ROOT = Path(__file__).resolve().parents[1]

VPDB_FASTA = ROOT / "Data" / "raw" / "vpdb" / "FungiDB-68_CaurisB8441_AnnotatedProteins.fasta"
UNIPROT_FILE = ROOT / "Data" / "processed" / "uniprot_clean.csv"
OUTPUT_FILE = ROOT / "Data" / "processed" / "vpdb_clean.csv"
MASTER_FILE = ROOT / "candida_master_mapping.xlsx"

HEADERS = [
    "VPDB Gene ID",
    "VPDB Product Name",
    "VPDB Sequence",
    "UniProt ID",
    "UniProt Name",
    "UniProt Sequence",
    "Sequence Identity %",
    "Sequence Match (Y/N)",
]


def parse_header(header):
    # pull the gene ID and product name from the FungiDB FASTA header
    gene = re.search(r"(?:^|\|\s*)gene=([^|]+)", header)
    product = re.search(r"(?:^|\|\s*)gene_product=([^|]+)", header)

    if gene is None:
        raise ValueError(f"gene ID missing from FASTA header: {header}")

    gene_id = gene.group(1).strip()
    product_name = product.group(1).strip() if product else ""

    return gene_id, product_name


def read_fasta(path):
    # read each FASTA record into a dataframe
    records = []
    header = None
    sequence = []

    with path.open(encoding="utf-8") as fasta:
        for raw_line in fasta:
            line = raw_line.strip()

            if not line:
                continue

            if line.startswith(">"):
                if header is not None:
                    gene_id, product_name = parse_header(header)
                    records.append({
                        "VPDB Gene ID": gene_id,
                        "VPDB Product Name": product_name,
                        "VPDB Sequence": "".join(sequence).upper(),
                    })

                header = line[1:]
                sequence = []
            else:
                sequence.append(line)

    # add the final FASTA record
    if header is not None:
        gene_id, product_name = parse_header(header)
        records.append({
            "VPDB Gene ID": gene_id,
            "VPDB Product Name": product_name,
            "VPDB Sequence": "".join(sequence).upper(),
        })

    if not records:
        raise ValueError("No proteins were parsed from the VPDB FASTA")

    return pd.DataFrame(records)


def add_uniprot_matches(vpdb):
    # exact sequence matching against the cleaned UniProt file
    uniprot = pd.read_csv(UNIPROT_FILE)

    needed = {"Entry", "Protein names", "Sequence"}
    missing = needed - set(uniprot.columns)

    if missing:
        raise KeyError(f"Missing UniProt columns: {', '.join(sorted(missing))}")

    uniprot["Sequence"] = (
        uniprot["Sequence"]
        .fillna("")
        .astype(str)
        .str.replace(r"\s+", "", regex=True)
        .str.upper()
    )

    # keep all accessions when the same sequence occurs more than once
    grouped = (
        uniprot[uniprot["Sequence"].ne("")]
        .groupby("Sequence", as_index=False)
        .agg({
            "Entry": lambda x: "; ".join(dict.fromkeys(x.dropna().astype(str))),
            "Protein names": lambda x: "; ".join(dict.fromkeys(x.dropna().astype(str))),
        })
        .rename(columns={
            "Sequence": "VPDB Sequence",
            "Entry": "UniProt ID",
            "Protein names": "UniProt Name",
        })
    )

    ambiguous = (
        uniprot[uniprot["Sequence"].ne("")]
        .groupby("Sequence")["Entry"]
        .nunique()
        .gt(1)
        .sum()
    )

    matched = vpdb.merge(grouped, on="VPDB Sequence", how="left")
    has_match = matched["UniProt ID"].notna()

    matched["UniProt Sequence"] = matched["VPDB Sequence"].where(has_match, "")
    matched["Sequence Identity %"] = has_match.map({True: 100.0, False: pd.NA})
    matched["Sequence Match (Y/N)"] = has_match.map({True: "Y", False: "N"})

    return matched[HEADERS], int(ambiguous)


def update_workbook(data):
    # write the mapping into the existing Master sheet
    workbook = load_workbook(MASTER_FILE)

    if "Master" not in workbook.sheetnames:
        raise KeyError("Master sheet not found in workbook")

    sheet = workbook["Master"]

    columns = {
        cell.value: cell.column
        for cell in sheet[1]
        if cell.value is not None
    }

    missing = [header for header in HEADERS if header not in columns]

    if missing:
        raise KeyError(f"Missing workbook headers: {', '.join(missing)}")

    # clear old mapping values before writing the new version
    for header in HEADERS:
        for row in range(2, sheet.max_row + 1):
            sheet.cell(row=row, column=columns[header]).value = None

    for row, record in enumerate(data.to_dict("records"), start=2):
        for header in HEADERS:
            value = record.get(header, "")

            sheet.cell(
                row=row,
                column=columns[header],
                value=None if pd.isna(value) else value,
            )

    workbook.save(MASTER_FILE)


def main():
    # check the files needed for the mapping are present
    for path in [VPDB_FASTA, UNIPROT_FILE, MASTER_FILE]:
        if not path.exists():
            raise FileNotFoundError(path)

    vpdb = read_fasta(VPDB_FASTA)
    vpdb, ambiguous = add_uniprot_matches(vpdb)

    # basic QC checks
    duplicates = vpdb["VPDB Gene ID"].duplicated().sum()
    empty_sequences = vpdb["VPDB Sequence"].eq("").sum()
    exact_matches = vpdb["Sequence Match (Y/N)"].eq("Y").sum()
    unmatched = vpdb["Sequence Match (Y/N)"].eq("N").sum()

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    vpdb.to_csv(OUTPUT_FILE, index=False)
    update_workbook(vpdb)

    # useful values to check after the mapping runs
    print(f"proteins: {len(vpdb)}")
    print(f"duplicate gene IDs: {duplicates}")
    print(f"empty sequences: {empty_sequences}")
    print(f"exact UniProt matches: {exact_matches}")
    print(f"unmatched proteins: {unmatched}")
    print(f"sequences with multiple UniProt accessions: {ambiguous}")


if __name__ == "__main__":
    main()


# GAI declaration
# GAI was used in this script for lines 7-130 and 133-195.
# i asked for a generated response on how to simplify the existing VPDB/UniProt
# mapping script while keeping the FASTA parsing, exact sequence matching and
# workbook update steps. I was given a shorter refactored version using regex
# header parsing, pandas groupby/merge operations and simplified workbook writing.
# that response was taken and used in this script.