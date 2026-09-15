from pathlib import Path
import re

import pandas as pd
from openpyxl import load_workbook

# file paths
ROOT = Path(__file__).resolve().parents[1]

GO_FILE = ROOT / "Data" / "raw" / "fungidb" / "GenesByTaxon_Summary.txt"
OUTPUT_FILE = ROOT / "Data" / "processed" / "go_terms_clean.csv"
MASTER_FILE = ROOT / "candida_master_mapping.xlsx"

COMPUTED_COLS = [
    "Computed GO Component IDs",
    "Computed GO Function IDs",
    "Computed GO Process IDs",
]
CURATED_COLS = [
    "Curated GO Component IDs",
    "Curated GO Function IDs",
    "Curated GO Process IDs",
]

GO_PATTERN = re.compile(r"GO:\d{7}")


def read_source(path):
    source = pd.read_csv(
        path,
        sep="\t",
        dtype=str,
        keep_default_na=False,
    )

    needed = ["Gene ID"] + COMPUTED_COLS + CURATED_COLS
    missing = [column for column in needed if column not in source.columns]

    if missing:
        raise KeyError(f"Missing source columns: {', '.join(missing)}")

    source["Gene ID"] = source["Gene ID"].str.strip()

    if source["Gene ID"].eq("").any():
        raise ValueError("Blank Gene IDs found in FungiDB source file")

    return source


def parse_terms(value):
    # split a GO cell and keep only valid GO accessions
    valid = []
    invalid = []

    for token in str(value).split(";"):
        token = token.strip()

        if not token or token.casefold() == "n/a":
            continue

        if GO_PATTERN.fullmatch(token):
            valid.append(token)
        else:
            invalid.append(token)

    return valid, invalid


def add_unique(target, values):
    # add terms without changing their original order
    for value in values:
        if value not in target:
            target.append(value)


def build_mapping(source):
    # combine computed and curated terms into one record per gene
    genes = {}
    invalid_tokens = []

    for record in source.to_dict("records"):
        gene_id = record["Gene ID"]

        if gene_id not in genes:
            genes[gene_id] = {
                "computed": [],
                "curated": [],
            }

        for column in COMPUTED_COLS:
            valid, invalid = parse_terms(record[column])
            add_unique(genes[gene_id]["computed"], valid)
            invalid_tokens.extend((gene_id, column, token) for token in invalid)

        for column in CURATED_COLS:
            valid, invalid = parse_terms(record[column])
            add_unique(genes[gene_id]["curated"], valid)
            invalid_tokens.extend((gene_id, column, token) for token in invalid)

    rows = []

    for gene_id, terms in genes.items():
        all_terms = terms["computed"].copy()
        add_unique(all_terms, terms["curated"])

        rows.append({
            "VPDB Gene ID": gene_id,
            "GO Terms": "; ".join(all_terms),
            "Has Computed GO": bool(terms["computed"]),
            "Has Curated GO": bool(terms["curated"]),
        })

    if not rows:
        raise ValueError("No GO records were created")

    return pd.DataFrame(rows), invalid_tokens


def update_workbook(go_terms):
    workbook = load_workbook(MASTER_FILE)

    if "Master" not in workbook.sheetnames:
        raise KeyError("Master sheet not found in workbook")

    sheet = workbook["Master"]
    columns = {
        cell.value: cell.column
        for cell in sheet[1]
        if cell.value is not None
    }

    for header in ["VPDB Gene ID", "GO Terms"]:
        if header not in columns:
            raise KeyError(f"Missing workbook header: {header}")

    # keep row numbers so GO terms can be written back to the same rows
    master = []
    for row in range(2, sheet.max_row + 1):
        gene_id = sheet.cell(row=row, column=columns["VPDB Gene ID"]).value

        if gene_id:
            master.append({
                "Workbook Row": row,
                "VPDB Gene ID": str(gene_id).strip(),
            })

    master = pd.DataFrame(master)
    merged = master.merge(go_terms, on="VPDB Gene ID", how="left")

    go_column = columns["GO Terms"]

    for row in range(2, sheet.max_row + 1):
        sheet.cell(row=row, column=go_column).value = None

    for record in merged.to_dict("records"):
        value = record.get("GO Terms")

        sheet.cell(
            row=record["Workbook Row"],
            column=go_column,
            value=None if pd.isna(value) or not str(value).strip() else value,
        )

    workbook.save(MASTER_FILE)

    source_ids = set(go_terms["VPDB Gene ID"])
    master_ids = set(master["VPDB Gene ID"])

    return {
        "mapped": merged["VPDB Gene ID"].isin(source_ids).sum(),
        "with_go": merged["GO Terms"].fillna("").str.strip().ne("").sum(),
        "source_only": len(source_ids - master_ids),
        "master_only": len(master_ids - source_ids),
    }


def main():
    for path in [GO_FILE, MASTER_FILE]:
        if not path.exists():
            raise FileNotFoundError(path)

    source = read_source(GO_FILE)
    go_terms, invalid = build_mapping(source)

    duplicate_ids = source["Gene ID"].duplicated().sum()
    computed = go_terms["Has Computed GO"].sum()
    curated = go_terms["Has Curated GO"].sum()
    both = (go_terms["Has Computed GO"] & go_terms["Has Curated GO"]).sum()

    unique_go = set()
    for value in go_terms["GO Terms"]:
        if value:
            unique_go.update(value.split("; "))

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    go_terms[["VPDB Gene ID", "GO Terms"]].to_csv(OUTPUT_FILE, index=False)

    stats = update_workbook(go_terms)

    print(f"genes: {len(go_terms)}")
    print(f"rows with GO terms: {stats['with_go']}")
    print(f"unique GO accessions: {len(unique_go)}")
    print(f"genes with computed GO: {computed}")
    print(f"genes with curated GO: {curated}")
    print(f"genes with both: {both}")
    print(f"duplicate source gene IDs: {duplicate_ids}")
    print(f"source IDs absent from master: {stats['source_only']}")
    print(f"master IDs absent from source: {stats['master_only']}")
    print(f"invalid GO tokens ignored: {len(invalid)}")


if __name__ == "__main__":
    main()


# GAI declaration
# GAI was used in this script for lines 50-66.
# i asked for a generated response on how to split semicolon-separated values,
# ignore blanks or "N/A", and check that each remaining value matched a GO
# accession such as GO:0008150. i was given a generic example using
# "your_value".split(";") together with re.fullmatch(r"GO:\d{7}", token).
# that was adapted in parse_terms() to separate valid and invalid GO terms.

# GAI was also used in this script for lines 76-116.
# i asked for a generated response on how to combine several annotation columns
# into one list per gene without keeping duplicate values, while still recording
# whether values came from two different groups of columns. i was given a generic
# example using a dictionary keyed by "your_id_column", looping over
# "your_group_1_columns" and "your_group_2_columns", then creating one output row
# per ID. that approach was adapted to combine computed and curated FungiDB GO
# annotations while retaining the Has Computed GO and Has Curated GO flags.