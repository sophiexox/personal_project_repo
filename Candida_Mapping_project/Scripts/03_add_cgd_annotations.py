from pathlib import Path
import re

import pandas as pd
from openpyxl import load_workbook

# file paths
ROOT = Path(__file__).resolve().parents[1]

CGD_FASTA = ROOT / "Data" / "raw" / "cgd" / "C_auris_B8441_current_orf_trans_all.fasta"
OUTPUT_FILE = ROOT / "Data" / "processed" / "cgd_clean.csv"
MASTER_FILE = ROOT / "candida_master_mapping.xlsx"

ANNOTATION_HEADERS = ["CGD Annotation", "CGD Name", "CDG Name"]


def parse_header(header):
    # extract the gene ID, CGD ID and annotation from each FASTA header
    gene = re.match(r"(\S+)", header)
    cgd_id = re.search(r"(?:^|\s)CGDID:(\S+)", header)
    annotation = re.search(r"\(\d+\s+amino acids\)\s*(.*)$", header)

    if not all([gene, cgd_id, annotation]):
        raise ValueError(f"Could not parse FASTA header: {header}")

    return (
        gene.group(1).strip(),
        cgd_id.group(1).strip(),
        annotation.group(1).strip(),
    )


def read_fasta(path):
    # only the FASTA headers are needed for this mapping
    records = []

    with path.open(encoding="utf-8") as fasta:
        for line in fasta:
            line = line.strip()

            if not line.startswith(">"):
                continue

            gene_id, cgd_id, annotation = parse_header(line[1:])
            records.append({
                "VPDB Gene ID": gene_id,
                "CGD ID": cgd_id,
                "CGD Annotation": annotation,
            })

    if not records:
        raise ValueError("No CGD records were parsed")

    return pd.DataFrame(records)


def join_unique(values):
    # keep unique IDs in their original order
    values = values.fillna("").astype(str).str.strip()
    return "; ".join(dict.fromkeys(values[values.ne("")]))


def longest(values):
    # keep the longest annotation if a gene has duplicate records
    values = values.fillna("").astype(str).str.strip()
    values = list(dict.fromkeys(values[values.ne("")]))
    return max(values, key=len, default="")


def collapse_duplicates(cgd):
    return (
        cgd.groupby("VPDB Gene ID", as_index=False, sort=False)
        .agg({
            "CGD ID": join_unique,
            "CGD Annotation": longest,
        })
    )


def update_workbook(cgd):
    workbook = load_workbook(MASTER_FILE)

    if "Master" not in workbook.sheetnames:
        raise KeyError("Master sheet not found in workbook")

    sheet = workbook["Master"]
    columns = {
        cell.value: cell.column
        for cell in sheet[1]
        if cell.value is not None
    }

    if "VPDB Gene ID" not in columns or "CGD ID" not in columns:
        raise KeyError("Required CGD columns are missing from the Master sheet")

    annotation_header = next(
        (name for name in ANNOTATION_HEADERS if name in columns),
        None,
    )

    if annotation_header is None:
        raise KeyError("No CGD annotation column found in the Master sheet")

    # read the existing gene IDs and workbook row numbers
    master = []
    for row in range(2, sheet.max_row + 1):
        gene_id = sheet.cell(row=row, column=columns["VPDB Gene ID"]).value

        if gene_id:
            master.append({
                "Workbook Row": row,
                "VPDB Gene ID": str(gene_id).strip(),
            })

    master = pd.DataFrame(master)

    if master.empty:
        raise ValueError("No VPDB Gene IDs found in the Master sheet")

    merged = master.merge(
        cgd,
        on="VPDB Gene ID",
        how="left",
        validate="many_to_one",
    )

    # clear old CGD values before writing the new mapping
    for header in ["CGD ID", annotation_header]:
        for row in range(2, sheet.max_row + 1):
            sheet.cell(row=row, column=columns[header]).value = None

    for record in merged.to_dict("records"):
        row = record["Workbook Row"]
        cgd_id = record.get("CGD ID")
        annotation = record.get("CGD Annotation")

        sheet.cell(
            row=row,
            column=columns["CGD ID"],
            value=None if pd.isna(cgd_id) else cgd_id,
        )
        sheet.cell(
            row=row,
            column=columns[annotation_header],
            value=None if pd.isna(annotation) else annotation,
        )

    workbook.save(MASTER_FILE)

    mapped = merged["CGD ID"].notna().sum()
    return int(mapped), len(merged) - int(mapped)


def main():
    for path in [CGD_FASTA, MASTER_FILE]:
        if not path.exists():
            raise FileNotFoundError(path)

    cgd_records = read_fasta(CGD_FASTA)

    # useful checks before duplicate records are collapsed
    duplicate_ids = cgd_records["VPDB Gene ID"].duplicated().sum()
    conflicting_ids = (
        cgd_records.groupby("VPDB Gene ID")["CGD ID"]
        .nunique()
        .gt(1)
        .sum()
    )
    empty_annotations = cgd_records["CGD Annotation"].eq("").sum()

    cgd = collapse_duplicates(cgd_records)

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    cgd.to_csv(OUTPUT_FILE, index=False)

    mapped, unmapped = update_workbook(cgd)

    print(f"records: {len(cgd_records)}")
    print(f"unique gene IDs: {len(cgd)}")
    print(f"duplicate gene IDs: {duplicate_ids}")
    print(f"conflicting CGD IDs: {conflicting_ids}")
    print(f"empty annotations: {empty_annotations}")
    print(f"mapped rows: {mapped}")
    print(f"unmapped rows: {unmapped}")


if __name__ == "__main__":
    main()


# GAI declaration
# GAI was used in this script for lines 17-30 and 57-77.
# i asked for a generated response on how to extract several pieces of information
# from structured FASTA headers and how to combine duplicate rows without losing
# unique identifiers. the response showed a generic example using regular
# expressions such as re.search() on "your_header" and pandas groupby().agg()
# on "your_dataframe". that approach was adapted to the CGD FASTA format and
# the VPDB Gene ID, CGD ID and CGD Annotation columns used in this script.