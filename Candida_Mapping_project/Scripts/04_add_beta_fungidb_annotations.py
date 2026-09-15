from pathlib import Path
import re
from urllib.parse import unquote

import pandas as pd
from openpyxl import load_workbook

# file paths
ROOT = Path(__file__).resolve().parents[1]

GFF_FILE = ROOT / "Data" / "raw" / "betafungidb" / "FungiDB-70_CaurisB8441.gff"
MAPPING_FILE = ROOT / "Data" / "raw" / "betafungidb" / "FungiDB-70_CaurisB8441_UniProtMapping.txt"
OUTPUT_FILE = ROOT / "Data" / "processed" / "betafungidb_clean.csv"
MASTER_FILE = ROOT / "candida_master_mapping.xlsx"

GENE_PATTERN = re.compile(r"B9J08_\d+")
HEADERS = ["VPDB Gene ID", "Beta FungiDB ID", "Beta FungiDB Name"]


def parse_attributes(text):
    # split the final GFF column into key/value pairs
    attributes = {}

    for field in text.split(";"):
        key, separator, value = field.partition("=")

        if separator:
            attributes[unquote(key.strip())] = unquote(value.strip())

    return attributes


def get_gene_id(attributes):
    # gene ID can appear under several different GFF fields
    for name in ["ID", "gene_id", "Name", "locus_tag"]:
        match = GENE_PATTERN.search(attributes.get(name, ""))

        if match:
            return match.group()

    for value in attributes.values():
        match = GENE_PATTERN.search(value)

        if match:
            return match.group()

    return ""


def get_description(attributes):
    for name in ["description", "product", "gene_product", "Note"]:
        value = attributes.get(name, "").strip()

        if value:
            return value

    return ""


def read_gff(path):
    # keep one record for each protein-coding gene
    records = []

    with path.open(encoding="utf-8") as gff:
        for line_number, line in enumerate(gff, start=1):
            line = line.rstrip()

            if not line or line.startswith("#"):
                continue

            fields = line.split("\t")

            if len(fields) != 9:
                raise ValueError(f"Invalid GFF row at line {line_number}")

            if fields[2] != "protein_coding_gene":
                continue

            attributes = parse_attributes(fields[8])
            gene_id = get_gene_id(attributes)

            if not gene_id:
                raise ValueError(f"No B9J08 ID found at line {line_number}")

            records.append({
                "Beta FungiDB ID": gene_id,
                "Beta FungiDB Name": get_description(attributes),
            })

    if not records:
        raise ValueError("No protein-coding genes were parsed from the GFF")

    return pd.DataFrame(records)


def longest(values):
    # use the longest non-empty description for duplicate gene records
    values = values.fillna("").astype(str).str.strip()
    values = list(dict.fromkeys(values[values.ne("")]))
    return max(values, key=len, default="")


def collapse_duplicates(data):
    return (
        data.groupby("Beta FungiDB ID", as_index=False, sort=False)
        .agg({"Beta FungiDB Name": longest})
    )


def read_mapping(path):
    # reference mapping supplied with FungiDB release 70
    mapping = pd.read_csv(
        path,
        sep="\t",
        header=None,
        names=["UniProt Accession", "Beta FungiDB ID", "FungiDB Gene URL"],
        dtype=str,
        comment="#",
    ).fillna("")

    for column in mapping.columns:
        mapping[column] = mapping[column].str.strip()

    invalid = ~mapping["Beta FungiDB ID"].str.fullmatch(GENE_PATTERN)

    if invalid.any():
        raise ValueError("Invalid B9J08 IDs found in UniProt mapping file")

    return mapping


def update_workbook(beta):
    workbook = load_workbook(MASTER_FILE)

    if "Master" not in workbook.sheetnames:
        raise KeyError("Master sheet not found in workbook")

    sheet = workbook["Master"]
    columns = {
        cell.value: cell.column
        for cell in sheet[1]
        if cell.value is not None
    }

    missing = [header for header in HEADERS if header not in columns]

    if missing:
        raise KeyError(f"Missing workbook headers: {', '.join(missing)}")

    # keep the workbook row number so merged values go back to the correct row
    master = []
    for row in range(2, sheet.max_row + 1):
        gene_id = sheet.cell(row=row, column=columns["VPDB Gene ID"]).value

        if gene_id:
            master.append({
                "Workbook Row": row,
                "VPDB Gene ID": str(gene_id).strip(),
            })

    master = pd.DataFrame(master)

    merged = master.merge(
        beta,
        left_on="VPDB Gene ID",
        right_on="Beta FungiDB ID",
        how="left",
        validate="many_to_one",
    )

    # clear old beta FungiDB values
    for header in ["Beta FungiDB ID", "Beta FungiDB Name"]:
        for row in range(2, sheet.max_row + 1):
            sheet.cell(row=row, column=columns[header]).value = None

    # write the updated mapping back into the same workbook rows
    for record in merged.to_dict("records"):
        row = record["Workbook Row"]

        for header in ["Beta FungiDB ID", "Beta FungiDB Name"]:
            value = record.get(header)

            sheet.cell(
                row=row,
                column=columns[header],
                value=None if pd.isna(value) else value,
            )

    workbook.save(MASTER_FILE)

    mapped = merged["Beta FungiDB ID"].notna().sum()
    release_ids = set(beta["Beta FungiDB ID"])
    master_ids = set(master["VPDB Gene ID"])

    return int(mapped), len(merged) - int(mapped), len(release_ids - master_ids)


def main():
    for path in [GFF_FILE, MAPPING_FILE, MASTER_FILE]:
        if not path.exists():
            raise FileNotFoundError(path)

    beta_records = read_gff(GFF_FILE)
    mapping = read_mapping(MAPPING_FILE)

    # checks before duplicate release records are collapsed
    duplicate_ids = beta_records["Beta FungiDB ID"].duplicated().sum()
    conflicting_names = (
        beta_records.groupby("Beta FungiDB ID")["Beta FungiDB Name"]
        .nunique()
        .gt(1)
        .sum()
    )
    missing_names = beta_records["Beta FungiDB Name"].fillna("").str.strip().eq("").sum()
    unspecified = (
        beta_records["Beta FungiDB Name"]
        .fillna("")
        .str.strip()
        .str.casefold()
        .eq("unspecified product")
        .sum()
    )

    beta = collapse_duplicates(beta_records)

    # compare release 70 genes against the supplied UniProt mapping
    release_ids = set(beta["Beta FungiDB ID"])
    reference_ids = set(mapping["Beta FungiDB ID"])

    missing_from_reference = len(release_ids - reference_ids)
    missing_from_release = len(reference_ids - release_ids)
    duplicate_reference_ids = mapping["Beta FungiDB ID"].duplicated().sum()

    url_mismatches = sum(
        not url.rstrip("/").endswith(f"/{gene_id}")
        for gene_id, url in zip(
            mapping["Beta FungiDB ID"],
            mapping["FungiDB Gene URL"],
        )
    )

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    beta.to_csv(OUTPUT_FILE, index=False)

    mapped, unmapped, absent_from_master = update_workbook(beta)

    print(f"release 70 genes: {len(beta)}")
    print(f"duplicate gene records: {duplicate_ids}")
    print(f"conflicting descriptions: {conflicting_names}")
    print(f"missing descriptions: {missing_names}")
    print(f"unspecified products: {unspecified}")
    print(f"mapped master rows: {mapped}")
    print(f"unmapped master rows: {unmapped}")
    print(f"release genes absent from master: {absent_from_master}")
    print(f"release genes absent from UniProt mapping: {missing_from_reference}")
    print(f"UniProt mapping genes absent from release: {missing_from_release}")
    print(f"duplicate mapping IDs: {duplicate_reference_ids}")
    print(f"gene URL/ID mismatches: {url_mismatches}")


if __name__ == "__main__":
    main()


# GAI declaration
# GAI was used in this script for lines 20-30.
# i asked for a generated response on how to parse the attributes column from a
# GFF3 file where each field is separated by a semicolon and some values are URL
# encoded. i was given a generic example using
# "your_attribute_text".split(";"), field.partition("=") and
# urllib.parse.unquote() to build a dictionary of attribute names and values.
# that was taken and adapted for the FungiDB GFF in parse_attributes().

# GAI was also used in this script for lines 150-189.
# i asked for a generated response on how to merge information from a pandas
# dataframe into an existing Excel workbook by matching an ID, while keeping the
# original workbook row order. i was given a generic example which stored
# "your_row_number" alongside "your_id_column", merged this dataframe with
# "your_dataframe" using pandas merge(), then used the stored row number when
# writing values back with openpyxl. that was taken and adapted to match
# Beta FungiDB IDs against VPDB Gene IDs in the Master worksheet.